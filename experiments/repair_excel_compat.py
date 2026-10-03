"""Create a conservative Excel-compatible workbook from the exported JSON."""
from pathlib import Path
import json
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
PAYLOAD = json.loads((ROOT / "Data/Expanded/intermediate/ivl_single_season_dataset.json").read_text(encoding="utf-8"))
OUTPUT = ROOT / "Data/Expanded/IVL_single_season_dataset_兼容版.xlsx"

NAVY = "253F63"
WHITE = "FFFFFF"
LIGHT = "DDEBF7"
header_fill = PatternFill("solid", fgColor=NAVY)
header_font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
body_font = Font(name="Microsoft YaHei", size=10)


def value(v, key):
    if v is None:
        return None
    if key in {"date", "日期"} and isinstance(v, str):
        try:
            return datetime.fromisoformat(v[:10])
        except ValueError:
            return v
    # Stop source text beginning with '=' from becoming a formula.
    if isinstance(v, str) and v.startswith("="):
        return "'" + v
    return v


def add_data_sheet(wb, title, rows, display=None, table_name="DataTable"):
    ws = wb.create_sheet(title)
    keys = list(rows[0])
    ws.append([display.get(k, k) if display else k for k in keys])
    for row in rows:
        ws.append([value(row.get(k), k) for k in keys])
    ws.freeze_panes = "B2" if title != "半局原始数据" else "D2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 34
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd"
    for idx, key in enumerate(keys, 1):
        width = 17
        if key == "小局唯一ID": width = 52
        elif key in {"大场对阵", "deciding_rule"}: width = 24
        elif key in {"date", "日期"}: width = 13
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = width
    table = Table(displayName=table_name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    return ws


wb = Workbook()
summary = wb.active
summary.title = "概览"
summary.sheet_view.showGridLines = False
summary["A2"] = "2025 IVL 秋季赛数据概览"
summary["A2"].font = Font(name="Microsoft YaHei", size=15, bold=True)
summary.append([])
summary.append(["指标", "值"])
m = PAYLOAD["manifest"]
match_counts = {t: sum(r["target"] == t for r in PAYLOAD["matches"]) for t in [1, 0, 0.5]}
game_counts = {t: sum(r["单局结果"] == t for r in PAYLOAD["small_games"]) for t in ["监管胜", "平局", "求生胜"]}
values = [
    ("半局记录", len(PAYLOAD["small_games"])), ("大场数", len(PAYLOAD["matches"])), ("队伍数", m["teams"]),
    ("主胜大场", match_counts[1]), ("客胜大场", match_counts[0]), ("平局或未决大场", match_counts[0.5]),
    ("半局监管胜", game_counts["监管胜"]), ("半局平局", game_counts["平局"]), ("半局求生胜", game_counts["求生胜"]),
    ("数据日期", f'{m["date_min"]} 至 {m["date_max"]}'), ("大场标签", "回合胜数优先，再比较总分"),
    ("口径差异场数", m["label_disagreements"]), ("来源", "Data/Data_details_All_games_details.csv")]
for item in values: summary.append(item)
for cell in summary[4]:
    cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center")
for row in summary.iter_rows(min_row=5, max_row=summary.max_row):
    for cell in row: cell.font = body_font
summary.column_dimensions["A"].width = 25
summary.column_dimensions["B"].width = 50

display = {"match_number":"大场序号","date":"日期","home":"主队","away":"客队","actual":"大场结果","elo_diff_scaled":"Elo差 / 400","smoothed_win_rate_diff":"历史胜率差","avg_point_margin_diff":"历史分差 / 10","recent5_margin_diff":"近五场分差 / 10","hunter_margin_diff":"监管分差 / 4","survivor_margin_diff":"求生分差 / 4","head_to_head_margin":"交锋分差 / 10"}
add_data_sheet(wb, "大场", PAYLOAD["matches"], table_name="MatchesTable")
add_data_sheet(wb, "赛前特征", PAYLOAD["features"], display, "FeaturesTable")
add_data_sheet(wb, "半局原始数据", PAYLOAD["small_games"], table_name="GamesTable")
wb.save(OUTPUT)

# Independent reopen verification.
check = load_workbook(OUTPUT, read_only=False, data_only=False)
assert check.sheetnames == ["概览", "大场", "赛前特征", "半局原始数据"]
assert check["大场"].max_row == 91
assert check["赛前特征"].max_row == 91
assert check["半局原始数据"].max_row == 487 and check["半局原始数据"].max_column == 99
assert check["概览"]["B5"].value == 486
check.close()
print(OUTPUT)
print("verified: sheets=4, matches=90, features=90, raw_rows=486, raw_columns=99")
